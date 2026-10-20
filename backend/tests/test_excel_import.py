from pathlib import Path

import pytest

from app.excel_import import TemplateError, parse_upload

FIXTURE_XLSX = Path(__file__).resolve().parent / "fixtures" / "dpwh_sample.xlsx"


@pytest.fixture()
def fixture_bytes():
    return FIXTURE_XLSX.read_bytes()


def test_fixture_parses_with_required_columns(fixture_bytes):
    result = parse_upload(fixture_bytes)
    assert result.total_rows_processed == 18


def test_fixture_ignores_stray_legend_column(fixture_bytes):
    result = parse_upload(fixture_bytes)
    assert "District Engineering Office (DEO)" in result.ignored_columns


def test_plain_meters_station(fixture_bytes):
    result = parse_upload(fixture_bytes)
    plan = next(p for p in result.accepted if p.section_id == "SEC_PLAIN")
    seg = plan.segments[0]
    assert seg.start_m == 0 and seg.start_text == "0"
    assert seg.end_m == 254 and seg.end_text == "254"


def test_chainage_text_station(fixture_bytes):
    result = parse_upload(fixture_bytes)
    plan = next(p for p in result.accepted if p.section_id == "SEC_CHAIN")
    seg = plan.segments[0]
    assert seg.start_m == 105
    assert seg.start_text == "Chainage 105"


def test_negative_offset_station(fixture_bytes):
    result = parse_upload(fixture_bytes)
    plan = next(p for p in result.accepted if p.section_id == "SEC_NEG")
    seg = plan.segments[0]
    assert seg.start_m == 5000 - 166
    assert seg.end_m == 5000
    assert seg.length_m == 166


def test_offset_over_1000_station(fixture_bytes):
    result = parse_upload(fixture_bytes)
    plan = next(p for p in result.accepted if p.section_id == "SEC_BIGOFFSET")
    seg = plan.segments[0]
    assert seg.end_m == 5000 + 1009


def test_decimal_offset_station(fixture_bytes):
    result = parse_upload(fixture_bytes)
    plan = next(p for p in result.accepted if p.section_id == "SEC_DECIMAL")
    seg = plan.segments[0]
    assert seg.start_m == pytest.approx(5099.80)
    assert seg.end_m == pytest.approx(5200.50)


def test_gap_is_accepted_with_warning(fixture_bytes):
    result = parse_upload(fixture_bytes)
    plan = next(p for p in result.accepted if p.section_id == "SEC_GAP")
    assert len(plan.segments) == 2
    assert any("Gap" in w.message for w in plan.warnings)


def test_overlap_is_accepted_with_warning(fixture_bytes):
    result = parse_upload(fixture_bytes)
    plan = next(p for p in result.accepted if p.section_id == "SEC_OVERLAP")
    assert any("overlaps" in w.message for w in plan.warnings)


def test_blank_shoulder_warns_but_literal_none_text_does_not(fixture_bytes):
    # SEC_BLANKSHOULDER: LEFT is a truly blank cell, RIGHT is the literal
    # text "None" -- both normalize to the stored value "none", but only
    # the truly blank cell is missing data worth warning about.
    result = parse_upload(fixture_bytes)
    plan = next(p for p in result.accepted if p.section_id == "SEC_BLANKSHOULDER")
    seg = plan.segments[0]
    assert seg.shoulder_left == "none"
    assert seg.shoulder_right == "none"
    assert len(plan.warnings) == 1
    assert plan.warnings[0].field == "LEFT"


def test_irregular_km_post_is_accepted_with_zero_length(fixture_bytes):
    result = parse_upload(fixture_bytes)
    plan = next(p for p in result.accepted if p.section_id == "SEC_IRREGULAR")
    seg = plan.segments[0]
    assert seg.irregular_km_post is True
    assert seg.length_m == 0
    assert plan.total_length_m == 0


def test_invalid_lanes_rejected(fixture_bytes):
    result = parse_upload(fixture_bytes)
    rejected_ids = {r.section_id for r in result.rejected}
    assert "SEC_BADLANES" in rejected_ids


def test_unknown_surface_rejected(fixture_bytes):
    result = parse_upload(fixture_bytes)
    rejected = next(r for r in result.rejected if r.section_id == "SEC_BADSURFACE")
    assert any("Invalid surface type" in e.message for e in rejected.errors)


def test_unparseable_station_rejected(fixture_bytes):
    result = parse_upload(fixture_bytes)
    rejected = next(r for r in result.rejected if r.section_id == "SEC_BADSTATION")
    assert any(e.field == "Start" for e in rejected.errors)


def test_multiple_deo_rejected(fixture_bytes):
    result = parse_upload(fixture_bytes)
    rejected = next(r for r in result.rejected if r.section_id == "SEC_MULTIDEO")
    assert any("multiple District Engineering Offices" in e.message for e in rejected.errors)


def test_blank_section_id_rejected(fixture_bytes):
    result = parse_upload(fixture_bytes)
    rejected_ids = {r.section_id for r in result.rejected}
    assert "(blank Section ID)" in rejected_ids


def test_classification_deo_ld_captured_on_plan(fixture_bytes):
    result = parse_upload(fixture_bytes)
    plan = next(p for p in result.accepted if p.section_id == "SEC_PLAIN")
    assert plan.classification == "Tertiary"
    assert plan.deo == "DEO A"
    assert plan.ld == "LD A"


def test_header_row_detected_when_not_row_two(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Some", "decorative", "banner", "row"])
    ws.append([None, None, None, None, None, "STATION"])
    ws.append(
        [
            "Section ID", "Classification", "DEO", "LD", "Road Name", "Start", "End",
            "Number of Lanes", "Surface type", "LEFT", "RIGHT",
        ]
    )
    ws.append(["S1", "Tertiary", "DEO A", "LD A", "Test Rd", "K0+000", "K1+000", 2, "Concrete", "Gravel", "Gravel"])
    out_path = tmp_path / "shifted_header.xlsx"
    wb.save(out_path)

    result = parse_upload(out_path.read_bytes())
    assert result.total_rows_processed == 1
    assert result.accepted[0].section_id == "S1"


def test_missing_required_column_raises_template_error(tmp_path):
    import openpyxl

    wb = openpyxl.load_workbook(FIXTURE_XLSX)
    ws = wb.active
    ws.cell(row=2, column=10).value = "Something Else"  # was 'LEFT'
    out_path = tmp_path / "broken.xlsx"
    wb.save(out_path)

    with pytest.raises(TemplateError, match="Missing required column"):
        parse_upload(out_path.read_bytes())
