import re
from dataclasses import asdict, dataclass, field
from io import BytesIO

import openpyxl

from app.chainage import ParsedRow, RowError, RowWarning, make_row, parse_station, validate_section_rows

# Maps normalized (lowercased, whitespace-collapsed) header text -> internal field name.
# Row 1 of the DPWH template is decorative group labels (STATION / SHOULDER);
# this maps the real column names on row 2. Column O ("District Engineering
# Office (DEO)" / "Legislative District (LD)" stray legend text) and any
# other unrecognized column are simply ignored.
REQUIRED_COLUMNS = {
    "section id": "section_id",
    "classification": "classification",
    "deo": "deo",
    "ld": "ld",
    "road name": "road_name",
    "start": "start_station",
    "end": "end_station",
    "number of lanes": "num_lanes",
    "surface type": "surface_type",
    "left": "shoulder_left",
    "right": "shoulder_right",
}

HEADER_SEARCH_MAX_ROW = 10  # how far down to look for the "Section ID" header row


class TemplateError(Exception):
    """Raised for file-level problems that abort the entire upload."""


def _normalize_header(name) -> str:
    return re.sub(r"\s+", " ", str(name if name is not None else "").strip().lower())


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _find_header_row(ws) -> int:
    """Finds the row holding the real column names by scanning for a cell
    that reads "Section ID" -- the DPWH template's row 1 is a merged
    group-header row (STATION / SHOULDER) above it, but that layout isn't
    guaranteed, so we detect the header row by content instead of a fixed
    row number."""
    for row_cells in ws.iter_rows(min_row=1, max_row=min(HEADER_SEARCH_MAX_ROW, ws.max_row)):
        for cell in row_cells:
            if _normalize_header(cell.value) == "section id":
                return cell.row
    raise TemplateError('Could not find the header row (a cell reading "Section ID") in the first '
                         f"{HEADER_SEARCH_MAX_ROW} rows")


@dataclass
class SectionImportPlan:
    section_id: str
    road_name: str
    classification: str
    deo: str
    ld: str
    segments: list[ParsedRow] = field(default_factory=list)
    warnings: list[RowWarning] = field(default_factory=list)

    @property
    def total_length_m(self) -> float:
        # Sum, not max(end_m): real sections aren't necessarily contiguous
        # from 0, and an irregular (length <= 0) segment contributes 0.
        return sum(s.length_m for s in self.segments)


@dataclass
class SectionRejection:
    section_id: str
    errors: list[RowError]


@dataclass
class ImportResult:
    accepted: list[SectionImportPlan]
    rejected: list[SectionRejection]
    total_rows_processed: int
    ignored_columns: list[str]


def import_result_to_dict(result: ImportResult) -> dict:
    """Plain-dict form of a parsed upload, JSON-serializable as-is (every
    field on every nested dataclass here is already a str/int/float/bool/
    list) -- used to persist a pending upload's parsed payload in the
    `pending_operations` table between preview and commit. See
    import_result_from_dict for the reverse; `asdict` alone can't round-trip
    back to the original dataclasses, only to this app's own reconstruction
    below."""
    return asdict(result)


def import_result_from_dict(d: dict) -> ImportResult:
    return ImportResult(
        accepted=[
            SectionImportPlan(
                section_id=p["section_id"],
                road_name=p["road_name"],
                classification=p["classification"],
                deo=p["deo"],
                ld=p["ld"],
                segments=[ParsedRow(**s) for s in p["segments"]],
                warnings=[RowWarning(**w) for w in p["warnings"]],
            )
            for p in d["accepted"]
        ],
        rejected=[
            SectionRejection(section_id=r["section_id"], errors=[RowError(**e) for e in r["errors"]])
            for r in d["rejected"]
        ],
        total_rows_processed=d["total_rows_processed"],
        ignored_columns=d["ignored_columns"],
    )


def parse_upload(file_bytes: bytes) -> ImportResult:
    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
        ws = wb.worksheets[0]
    except Exception as exc:  # noqa: BLE001 - surfaced as a template error to the client
        raise TemplateError(f"Could not read Excel file: {exc}") from exc

    header_row = _find_header_row(ws)
    header_cells = next(ws.iter_rows(min_row=header_row, max_row=header_row))
    found_headers = [c.value for c in header_cells if c.value is not None]
    normalized_to_col_idx: dict[str, int] = {}
    for c in header_cells:
        if c.value is not None:
            normalized_to_col_idx[_normalize_header(c.value)] = c.column

    missing = [req for req in REQUIRED_COLUMNS if req not in normalized_to_col_idx]
    if missing:
        missing_display = [c.upper() if c in ("left", "right", "deo", "ld") else c.title() for c in missing]
        raise TemplateError(
            f"Missing required column(s): {', '.join(missing_display)}. "
            f"Columns found in file: {found_headers}"
        )

    col_idx = {
        field_name: normalized_to_col_idx[norm] for norm, field_name in REQUIRED_COLUMNS.items()
    }
    used_col_indices = set(col_idx.values())
    ignored_columns = [
        str(c.value) for c in header_cells if c.value is not None and c.column not in used_col_indices
    ]

    rows_by_section: dict[str, list[ParsedRow]] = {}
    errors_by_section: dict[str, list[RowError]] = {}
    total_rows = 0

    def cell(row_cells, field_name):
        return row_cells[col_idx[field_name] - 1].value

    def text(row_cells, field_name) -> str:
        value = cell(row_cells, field_name)
        return "" if _is_blank(value) else str(value).strip()

    for row_cells in ws.iter_rows(min_row=header_row + 1):
        row_number = row_cells[0].row
        if all(c.value is None for c in row_cells):
            continue
        total_rows += 1

        section_id = text(row_cells, "section_id")
        if not section_id:
            errors_by_section.setdefault("(blank Section ID)", []).append(
                RowError(row_number, "Section ID", "Section ID is blank")
            )
            continue

        row_errors = errors_by_section.setdefault(section_id, [])

        start_raw = cell(row_cells, "start_station")
        end_raw = cell(row_cells, "end_station")
        try:
            start = parse_station(start_raw)
        except ValueError as exc:
            row_errors.append(RowError(row_number, "Start", str(exc)))
            continue
        try:
            end = parse_station(end_raw)
        except ValueError as exc:
            row_errors.append(RowError(row_number, "End", str(exc)))
            continue

        num_lanes_raw = cell(row_cells, "num_lanes")
        try:
            num_lanes = int(num_lanes_raw)
        except (TypeError, ValueError):
            row_errors.append(
                RowError(row_number, "Number of Lanes", f"'{num_lanes_raw}' is not a valid integer")
            )
            continue

        def _norm_surface(value) -> str:
            return "none" if _is_blank(value) else str(value).strip().lower()

        parsed = make_row(
            row_number=row_number,
            section_id=section_id,
            road_name=text(row_cells, "road_name"),
            classification=text(row_cells, "classification"),
            deo=text(row_cells, "deo"),
            ld=text(row_cells, "ld"),
            start=start,
            end=end,
            num_lanes=num_lanes,
            surface_type=_norm_surface(cell(row_cells, "surface_type")),
            shoulder_left=_norm_surface(cell(row_cells, "shoulder_left")),
            shoulder_right=_norm_surface(cell(row_cells, "shoulder_right")),
            # Only a truly empty cell counts as "blank" -- the literal text
            # "None" is a valid DPWH value (explicitly "no shoulder"), not
            # missing data, even though both normalize to "none" above.
            shoulder_left_was_blank=_is_blank(cell(row_cells, "shoulder_left")),
            shoulder_right_was_blank=_is_blank(cell(row_cells, "shoulder_right")),
        )
        rows_by_section.setdefault(section_id, []).append(parsed)

    accepted: list[SectionImportPlan] = []
    rejected: list[SectionRejection] = []

    all_section_ids = set(rows_by_section) | set(errors_by_section)
    for section_id in all_section_ids:
        rows = rows_by_section.get(section_id, [])
        pre_errors = list(errors_by_section.get(section_id, []))
        validation_errors, warnings = validate_section_rows(rows) if rows else ([], [])
        all_errors = pre_errors + validation_errors

        if all_errors:
            rejected.append(SectionRejection(section_id=section_id, errors=all_errors))
            continue

        sorted_rows = sorted(rows, key=lambda r: r.start_m)
        road_name = next((r.road_name for r in sorted_rows if r.road_name), section_id)
        classification = next((r.classification for r in sorted_rows if r.classification), "")
        deo = next((r.deo for r in sorted_rows if r.deo), "")
        ld = next((r.ld for r in sorted_rows if r.ld), "")
        accepted.append(
            SectionImportPlan(
                section_id=section_id,
                road_name=road_name,
                classification=classification,
                deo=deo,
                ld=ld,
                segments=sorted_rows,
                warnings=warnings,
            )
        )

    return ImportResult(
        accepted=accepted,
        rejected=rejected,
        total_rows_processed=total_rows,
        ignored_columns=ignored_columns,
    )
